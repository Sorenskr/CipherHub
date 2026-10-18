# -*- coding: utf-8 -*-
"""
从两个 Excel 表格导入初始数据，生成 data/db.json
用法: python tools/import_excel.py
"""
import os
import sys
import json
import hashlib
import secrets

BASE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(BASE, "data", "db.json")

FILE1 = r"C:\Users\OMEN 11\AppData\Local\Temp\codebuddy-dropped-files\705a4021-da7f-4b48-9a57-a43e6926c188\laya暗号和短链.xlsx"
FILE2 = r"C:\Users\OMEN 11\AppData\Local\Temp\codebuddy-dropped-files\ce9e77fa-1304-45d5-81b3-9c3bb172308c\Sonka暗号表.xls"

PLATFORMS = [
    {"id": "threads", "name": "Threads", "color": "#6e56cf", "icon": "threads", "order": 1},
    {"id": "instagram", "name": "Instagram", "color": "#e1306c", "icon": "instagram", "order": 2},
    {"id": "tiktok", "name": "TikTok", "color": "#00c8b4", "icon": "tiktok", "order": 3},
    {"id": "x", "name": "X", "color": "#8b93a7", "icon": "x", "order": 4},
    {"id": "facebook", "name": "Facebook", "color": "#1877f2", "icon": "facebook", "order": 5},
    {"id": "video", "name": "视频", "color": "#ff7a59", "icon": "video", "order": 6},
    {"id": "link", "name": "短链", "color": "#3ea6ff", "icon": "link", "order": 7},
    {"id": "image", "name": "图文", "color": "#ffb020", "icon": "image", "order": 8},
]

PEOPLE = {}
PEOPLE_ORDER = []


def s(v):
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def get_person(name, group):
    if not name:
        return None
    if name not in PEOPLE:
        PEOPLE[name] = {
            "id": len(PEOPLE) + 1,
            "name": name,
            "group": group or "",
            "note": "",
        }
        PEOPLE_ORDER.append(name)
    else:
        cur = PEOPLE[name]["group"]
        parts = [p for p in cur.split(" / ") if p]
        if group and group not in parts:
            parts.append(group)
            PEOPLE[name]["group"] = " / ".join(parts)
    return PEOPLE[name]


entries = []


def add_entry(person_name, platform_id, value, kind="code", type_="", url=""):
    if not value:
        return
    entries.append({
        "id": len(entries) + 1,
        "personId": PEOPLE[person_name]["id"] if person_name in PEOPLE else None,
        "platformId": platform_id,
        "kind": kind,
        "value": value,
        "type": type_,
        "url": url,
        "note": "",
    })


# ---------------- 文件1: laya 暗号和短链 ----------------
import openpyxl

wb = openpyxl.load_workbook(FILE1, data_only=True)

CODE_SHEETS = {
    "thresds": "threads",
    "ins": "instagram",
    "tk": "tiktok",
    "x": "x",
    "fb": "facebook",
    "视频": "video",
}

for title, pid in CODE_SHEETS.items():
    ws = wb[title]
    rows = list(ws.iter_rows(values_only=True))
    for r in rows[1:]:
        name = s(r[0])
        code = s(r[1])
        if not name or not code:
            continue
        get_person(name, "Laya")
        add_entry(name, pid, code, "code", "")

# 短链
ws = wb["短链"]
for r in list(ws.iter_rows(values_only=True))[1:]:
    name, slug, link, typ = s(r[0]), s(r[1]), s(r[2]), s(r[3])
    if not name or not slug:
        continue
    get_person(name, "Laya")
    add_entry(name, "link", slug, "link", typ, link)

# 长沙短链（公共，挂在虚拟成员"长沙公共"下）
CS_NAME = "长沙公共"
ws = wb["长沙短链"]
for r in list(ws.iter_rows(values_only=True))[1:]:
    slug, link, typ = s(r[0]), s(r[1]), s(r[2])
    if not slug:
        continue
    get_person(CS_NAME, "长沙")
    add_entry(CS_NAME, "link", slug, "link", typ, link)

# 长沙暗号（图文 / 视频 两列）
ws = wb["长沙暗号"]
rows = list(ws.iter_rows(values_only=True))
header = [s(c) for c in rows[0]]
col_img = header.index("图文") if "图文" in header else 0
col_vid = header.index("视频") if "视频" in header else 1
for r in rows[1:]:
    v1, v2 = s(r[col_img]), s(r[col_vid])
    if v1:
        get_person(CS_NAME, "长沙")
        add_entry(CS_NAME, "image", v1, "code", "图文")
    if v2:
        get_person(CS_NAME, "长沙")
        add_entry(CS_NAME, "video", v2, "code", "视频")

# ---------------- 文件2: Sonka 暗号表 ----------------
import xlrd

PLAT_MAP = {
    "thareads": "threads", "threads": "threads", "thresds": "threads",
    "ins": "instagram", "instagram": "instagram",
    "tik": "tiktok", "tiktok": "tiktok",
    "facebook": "facebook", "fb": "facebook",
    "x": "x", "X": "x",
}

bk = xlrd.open_workbook(FILE2)
sh = bk.sheet_by_index(0)
head = [s(sh.cell_value(0, c)) for c in range(sh.ncols)]
name_cols = [c for c, v in enumerate(head) if v == "姓名"]
for c in name_cols:
    if c + 2 >= sh.ncols:
        continue
    for r in range(1, sh.nrows):
        name = s(sh.cell_value(r, c))
        code = s(sh.cell_value(r, c + 1))
        plat = s(sh.cell_value(r, c + 2))
        if not name or not code:
            continue
        pid = PLAT_MAP.get(plat.lower(), "")
        if not pid:
            continue
        get_person(name, "Sonka")
        add_entry(name, pid, code, "code", "")

# ---------------- 输出 ----------------
salt = secrets.token_hex(16)
pass_hash = hashlib.sha256((salt + "admin123").encode("utf-8")).hexdigest()

people_list = [PEOPLE[n] for n in PEOPLE_ORDER]

db = {
    "version": 1,
    "settings": {
        "siteName": "暗号查询中心",
        "subtitle": "输入姓名或暗号，秒查全平台归属",
        "adminUser": "admin",
        "adminSalt": salt,
        "adminPassHash": pass_hash,
        "updatedAt": "",
    },
    "platforms": PLATFORMS,
    "groups": ["Laya", "Sonka", "长沙"],
    "people": people_list,
    "entries": entries,
}

os.makedirs(os.path.dirname(OUT), exist_ok=True)
with open(OUT, "w", encoding="utf-8") as f:
    json.dump(db, f, ensure_ascii=False, indent=1)

print("people:", len(people_list))
print("entries:", len(entries))
for p in PLATFORMS:
    n = sum(1 for e in entries if e["platformId"] == p["id"])
    print("  ", p["name"], n)
print("out:", OUT)
